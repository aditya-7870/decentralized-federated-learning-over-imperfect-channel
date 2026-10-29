"""Build editable academic report from completed results with bundled Python."""
import json
import sys
import re
import html
from report_sections import introduction, literature, approach, discussion
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/report';OUT.mkdir(parents=True,exist_ok=True)
summary=json.loads((ROOT/'results/review/summary.json').read_text(encoding='utf-8'))
groups=summary['groups']
spectra=json.loads((ROOT/'results/review/figures/spectral_table.json').read_text())
if '--preview' not in sys.argv and any(g['seeds']!=[42,43,44,45,46] for g in groups):raise RuntimeError('Complete the five specified seeds before building report')
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
section.top_margin=Inches(.55);section.bottom_margin=Inches(.6)
section.left_margin=section.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(11)
 s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.08
for name,size in [('Title',16),('Heading 1',13),('Heading 2',12),('Heading 3',11)]:
 doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
 doc.styles[name].paragraph_format.keep_with_next=True
 doc.styles[name].paragraph_format.space_before=Pt(9)
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
doc.styles['Caption'].font.size=Pt(10);doc.styles['Caption'].font.italic=True
doc.styles['Caption'].font.bold=False
# Remove inherited theme fonts and title rules from the bundled starter.
for style in doc.styles:
 for node in list(style.element.iter()):
  if node.tag==qn('w:pBdr'):
   node.getparent().remove(node)
  if node.tag==qn('w:rFonts'):
   for attr in list(node.attrib):
    if 'theme' in attr.lower():del node.attrib[attr]

doc.core_properties.author='Aditya Kumar Sahu; Rishav Kumar Manjhi; D. Yaswanth Reddy'
doc.core_properties.title='Communication Tradeoffs in Decentralized Learning over Lossy Links'
references=json.loads((ROOT/'report_sources/references.json').read_text(encoding='utf-8'))
reference_numbers={}
legacy_reference_keys={1:'mcmahan',2:'r3',3:'r1',4:'r2',5:'r5',6:'r14',7:'r10',8:'r11',9:'mnist',10:'cifar'}
def reference_text(text):
 text=str(text)
 if '{ref:' not in text:
  def legacy(m):
   ids=[int(v) for v in m.group(1).split(',')]
   return ''.join('{ref:'+legacy_reference_keys[i]+'}' for i in ids) if all(i in legacy_reference_keys for i in ids) else m.group(0)
  text=re.sub(r'\[([0-9]+(?:,[0-9]+)*)\]',legacy,text)
 def resolve(m):
  key=m.group(1)
  if key not in references:raise ValueError('Unknown reference '+key)
  if key not in reference_numbers:reference_numbers[key]=len(reference_numbers)+1
  return '['+str(reference_numbers[key])+']'
 return re.sub(r'\]\[','], [',re.sub(r'\{ref:([a-z0-9]+)\}',resolve,text))
def p(text):return doc.add_paragraph(reference_text(text))
def h(text):
 level=3 if re.match(r'^\d+\.\d+\.\d+',text) or (text!='Abstract' and not re.match(r'^\d+\.\d+',text)) else 2
 return doc.add_heading(text,level)
def page(title):
 doc.add_page_break()
 if title.endswith('continued'):
  para=doc.add_paragraph(title)
  para.paragraph_format.space_after=Pt(8)
  for run in para.runs:run.bold=True;run.italic=True;run.font.size=Pt(10.5)
 else:
  doc.add_heading(title,2 if re.match(r'^\d+\.\d+',title) else 1)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 if widths is None:widths=[6.96/len(headers)]*len(headers)
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=reference_text(v)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=reference_text(v)
 for ri,row in enumerate(t.rows):
  props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
  if ri==0:props.append(OxmlElement('w:tblHeader'))
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i])
   for para in c.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=1
    for run in para.runs:run.font.size=Pt(9.5);run.bold=(ri==0)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
table_count=0
def caption(text):
 global table_count
 if text.startswith('Table '):
  table_count+=1
  text=re.sub(r'^Table [0-9]+',f'Table {table_count}',text)
 doc.add_paragraph(reference_text(text),style='Caption')
def fig(name,text,width=6.85):
 path=ROOT/'results/review'/name
 if not path.exists():return
 para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
 para.add_run().add_picture(str(path),width=Inches(width));caption(text)
def eq(text,n):
 para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 math=OxmlElement('m:oMath');run=OxmlElement('m:r');txt=OxmlElement('m:t');txt.text=text;run.append(txt);math.append(run);para._p.append(math);para.add_run(f'    ({n})')
def select(**kw):return [g for g in groups if all(g.get(k)==v for k,v in kw.items())]
def base(partition,method,topology='ring',clients=5,dataset='MNIST'):
 return select(dataset=dataset,partition=partition,method=method,topology=topology,clients=clients,loss=0.0,delay=0,delta=0.0,rounds=10,local_epochs=1,batch_size=128)
def stat(g,key='final_accuracy',percent=True):
 a=g[key];scale=100 if percent else 1;half=a['ci95_half'] or 0
 precision=3 if 0 < half*scale < .005 else 2
 return f"{a['mean']*scale:.{precision}f} ± {half*scale:.{precision}f}"

para=doc.add_paragraph('COMMUNICATION TRADEOFFS IN DECENTRALIZED LEARNING OVER LOSSY LINKS',style='Title');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
for line in ['Aditya Kumar Sahu, Rishav Kumar Manjhi, D. Yaswanth Reddy','Department of Computer Science and Engineering, Vellore Institute of Technology (VIT), Vellore,','632014, Tamil Nadu, India','Corresponding author: aditya.kumarsahu2024@vitstudent.ac.in']:
 para=p(line);para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.space_after=Pt(3)
h('Abstract')
iid=base('iid','local')[0];two=base('two-class','local')[0];ring=base('two-class','gossip')[0]
p(f'Decentralized learning exchanges model parameters between neighboring clients without a central aggregation server. We evaluate when communication improves learning and what it costs under channel impairments. A PyTorch simulator compares local-only, standard gossip, and event-triggered learning using IID, exactly two-class, and Dirichlet label partitions. The core MNIST experiments use five independent seeds, ten gossip rounds, one local epoch per round, and a 206,922-parameter CNN. IID local-only accuracy is {stat(iid)}%, while two-class local-only accuracy is {stat(two)}%; Ring gossip reaches {stat(ring)}% on the same two-class setting. These controlled comparisons explain why high IID MNIST accuracy alone is insufficient evidence of channel robustness. Further experiments evaluate delay, independent and burst erasures, network size, and CIFAR-10. Traffic reduction is measured from actual directed messages and payload bytes. Spectral analysis verifies fixed-matrix averaging separately from local SGD. The conclusions characterize the tested configurations, rather than a universal topology ranking or a convergence theorem for impaired training.')
p('Keywords: decentralized learning; gossip; non-IID data; event-triggered communication; packet loss; delay; spectral gap; reproducibility.')
introduction(doc,p,h,page)
literature(p,h,page,table,caption)

page('3 Methodology')
h('3.1 Problem Formulation')
p('Client i holds a disjoint dataset Dᵢ with local mean cross-entropy loss Fᵢ. The intended objective weights clients uniformly. With unequal Dirichlet partition sizes, this differs from a sample-weighted pooled-data objective. Within each seed, all compared methods use identical partitions, reference weights, local hyperparameters, and per-client minibatch orders.')
eq('min F(w) = N⁻¹ Σᵢ Fᵢ(w)',1)
p('Core experiments use one local SGD epoch per round; the separate original-budget control uses two. Learning rate is 0.01 and momentum is 0.9. Only model parameters are mixed; optimizer momentum remains local. This policy is shared by all methods. The core budget is ten local epochs per client in total, while the original-budget control uses forty.')
h('Network Topologies')
p('Ring is a cycle on N nodes. Grid keeps the first N row-major nodes of a rectangular grid whose row count is the nearest integer to √N, with enough columns to contain N nodes. Random uses connected Erdős-Rényi G(N,0.5) samples. Graphs are static and undirected; every raw run records its actual edges. Experimental seeds vary the Random realization, while the illustrative five-node spectral table uses seed 42.')
eq('Wᵢⱼ = 1 / [1 + max(deg(i), deg(j))],  (i,j) ∈ E',2)
eq('Wᵢᵢ = 1 − Σⱼ≠ᵢ Wᵢⱼ',3)
p('For an undirected graph, the original Metropolis-Hastings matrix is symmetric and doubly stochastic. Let Rᵢ(t) contain the neighbors whose newest acceptable messages arrive during round t. The receiver combines its own post-training model with those arrivals:')
eq('wᵢ⁺ = [Wᵢᵢ wᵢ + Σⱼ∈Rᵢ(t) Wᵢⱼ wⱼ(received)] / [Wᵢᵢ + Σⱼ∈Rᵢ(t) Wᵢⱼ]',4)
p('With no acceptable arrival, the client retains its local model. Independently missing directed messages preserve row normalization but generally break column normalization. The arithmetic network mean can drift during aggregation. Consequently this policy is not described as unbiased global averaging under loss.')
fig('figures/architecture.png','Figure 1. Common training workflow. M3 analyzes the ideal graph separately from neural-network learning.')

page('3.1 Problem Formulation continued')
h('Channel Model')
h('Loss and burst model')
p('The independent-loss channel erases each directed message with probability 0.2 in impaired comparisons. Randomness is keyed by seed, sender, receiver, round, and stream. Suppressing an unrelated transmission therefore cannot shift the channel outcome associated with a given link and round.')
p('The Gilbert-Elliott model has good-to-bad transition probability 0.05 and bad-to-good probability 0.20 per round. Good-state erasure probability is zero; bad-state erasure probability is one. Stationary erasure probability is 0.05/(0.05+0.20)=0.20, matching the nominal independent-loss rate. Mean bad-state duration is five rounds. Each link starts in stationarity and evolves through silent rounds. Short-run realized losses remain variable and are retained in the raw counts.')
h('Delay and receiver freshness')
p('A surviving message receives an independent uniform integer delay between zero and dₘₐₓ inclusive. Maximum delays 1, 2, and 4 are compared with zero delay under 20% Bernoulli loss. Source-round metadata ensures that the newest arrival from a sender is selected, and an older version cannot replace one already accepted. Delivered messages discarded as stale still consume communication.')
p('Packets scheduled after the final round remain pending. Attempted, dropped, arrived, and pending messages are separate counts. Delivered bytes refer to arrivals within the observation horizon, not future scheduled deliveries. No post-training buffer drain changes the final models.')
h('M2 Event Triggered Gossip')
eq('sendᵢ(t) ⇔ δ = 0 or ||wᵢ(t) − wᵢ(last sent)||₂ > δ',5)
p('A triggered client sends to every neighbor and updates its last-sent snapshot even if the channel later erases the update. Silent clients can still receive and average. Threshold zero explicitly means always transmit. Revised thresholds 0.5 and 2.0 are fixed comparison settings, not thresholds selected by optimizing test accuracy.')
eq('message savings (%) = 100 [1 − actual attempts / (2 |E| T)]',6)
p('Broadcast suppression uses N T opportunities. Message savings weight a silent client by its degree. The percentages coincide on a regular Ring but can differ on Grid or Random graphs. Payload bytes equal attempted messages times the float32 model size; network-protocol overhead is excluded.')

page('3.2 Data and Materials')
h('Datasets and Client Partitions')
p('MNIST contains 60,000 training and 10,000 test images of handwritten digits. Grayscale 28×28 inputs are divided by 255 and normalized using mean 0.1307 and standard deviation 0.3081. IID allocation shuffles and divides indices. The two-class partition gives each client exactly two labels and splits each label among its assigned clients. Dirichlet allocation uses class-wise α=0.1 proportions and rejects a split with fewer than 100 samples at any client. Every split covers the full training set exactly once. Client sizes, class histograms, and partition hashes are saved.')
h('CNN Architecture')
table(['Layer','MNIST configuration','Output'],[
 ['Convolution, ReLU, max pool','1→16; 3×3; padding 1; pool 2','16×14×14'],['Convolution, ReLU, max pool','16→32; 3×3; padding 1; pool 2','32×7×7'],['Flatten and FC, ReLU','1568→128','128'],['Output logits','128→10','10'],['Trainable parameters','206,922','827,688 bytes/message']],[2.15,3.0,1.81])
caption('Table 2. Corrected MNIST CNN architecture and float32 communication payload.')
h('CIFAR 10 Materials')
p('CIFAR-10 adds 50,000 training and 10,000 test RGB images of size 32×32. The corresponding CNN uses three input channels and a 2048→128 first fully connected layer, giving 268,650 parameters. The extension applies the same Dirichlet α=0.1 allocation and compares local-only with Ring gossip. MNIST and CIFAR-10 are public benchmark datasets {ref:mnist}{ref:cifar}.')
h('Implementation Materials')
p('The simulator uses Python and PyTorch for local learning, NumPy and NetworkX for numerical and graph operations, SciPy for Student-t intervals, and Matplotlib for figures. Data and model tensors are held on the training device to avoid repeated per-image preprocessing. Models are saved in a format that can be loaded onto a CPU for review.')
approach(p,h,page)
page('3.4 Experimental Setup')
table(['Setting','Revised value'],[
 ['Independent seeds','42, 43, 44, 45, 46'],['Rounds / local epochs','10 / 1'],['Batch / learning rate / momentum','128 / 0.01 / 0.9'],['Core size and partitions','5 clients; IID, two-class, Dirichlet α=0.1'],['Scale comparison','20 and 50 clients; two-class; Ring and Grid'],['Loss comparison','20% Bernoulli; 20% stationary Gilbert-Elliott'],['Delay comparison','Maximum delay 0, 1, 2, 4 rounds'],['Event comparison','Dirichlet; Ring and Grid; δ=0, 0.5, 2.0']],[2.65,4.31])
caption('Table 3. Revised experimental protocol. The core budget is ten rounds with one local epoch. A separate matched IID control uses twenty rounds and two local epochs.')
p('All client models are evaluated on the full shared test set. Reported accuracy is their arithmetic mean, not an ensemble or centrally averaged predictor. Training uses the detected RTX 5070 Ti Laptop GPU and CUDA-enabled PyTorch with automatic mixed precision for computation. Exchanged parameters remain float32. Deterministic algorithms are enabled, and each run records software versions, device, graph edges, and complete configuration.')

h('Experiment Families and Evaluation Scope')
table(['Experiment family','Conditions','Seeds each','Runs'],[
 ['Core MNIST: three partitions × local/Ring/Grid/Random',12,5,60],
 ['Event thresholds: two topologies × two nonzero thresholds',4,5,20],
 ['Independent/burst erasures and finite delays',7,5,35],
 ['Scale: two client counts × Ring/Grid',4,5,20],
 ['Original-budget IID: local-only/Ring',2,5,10],
 ['CIFAR-10 Dirichlet: local-only/Ring',2,5,10],
 ['Total reported experiments',31,'5',155]],[3.9,1.0,1.06,1.0])
caption('Table 5. Completed experimental design. Five additional ten-round GPU runs check M1/M2 zero-threshold equivalence and are not counted as independent comparison conditions.')
p('The comparisons are controlled experiment families rather than a full cross-product of every setting. Delay and burst-loss learning use five-client Ring with Dirichlet MNIST; independent zero-delay loss additionally compares Grid and Random. Scale uses two-class MNIST on Ring and Grid. The original-budget IID control uses twenty rounds, two local epochs, and batch size 64. CIFAR-10 evaluates the restricted comparison described above.')

page('3.5 Evaluation and Reproducibility')
eq('E(t) = N⁻¹ Σᵢ ||wᵢ(t) − w̄(t)||₂²',7)
p('Consensus error measures squared parameter disagreement. It is not a substitute for test accuracy: neural-network symmetries and repeated local updates can change the relationship. Aggregation-only mean drift is the norm of the change in the network mean immediately across averaging, excluding the preceding local training step.')
p('Every reported condition uses five independent seed-level results. The 95% interval is mean ± t₀.₉₇₅,ₙ₋₁ s/√n. It summarizes variation from initialization, data allocation, minibatches, graph sampling where applicable, and channel realization. It does not measure generalization across entirely different datasets or algorithm families. Five seeds provide limited precision; overlapping intervals are not used as a significance test.')
p('Gossip-minus-local accuracy gains are also computed within each matched seed and then summarized with a Student-t interval. This paired estimate is preferable to subtracting independently aggregated interval endpoints. Conditions with different client counts are not pooled, and thresholds are not chosen retrospectively from the test results.')
h('Validation checks')
p('Full-budget GPU validation repeats ten-round Dirichlet training across all five seeds. Standard gossip and zero-threshold M2 match exact initial and final client-model hashes and per-round mean accuracy in the zero-loss Ring setting. Regression tests compare both training entry points on synthetic data under loss and delay. Additional checks cover pending arrival counts, randomness under suppressed sends, stale-version rejection, unequal-degree message savings, signed eigenvalues, parameter count, the squared-error averaging bound, partition coverage, and average drift under one-way reception.')
h('Reproducible Artifacts and Data Availability')
p('Revised raw runs are in results/review/runs. Aggregate statistics, plots, a validation record, and a legacy accounting audit accompany them. Filenames identify the condition, seed, and training budget. Completed results are saved atomically; the runner skips completed files when restarted. Partially trained runs restart from their initial state. Original reports and original result JSON files remain unchanged.')
p('The portable CPU demonstration includes a trained MNIST model and twenty bundled test digits. It requires no retraining or dataset download. A separate standard-library command prints the saved experiment summary. Dependency specifications and handover commands allow review on a laptop without the training GPU. Source code, raw results, aggregate statistics, a dependency specification, and a portable checkpoint accompany the report.')

page('4 Results')
h('4.1 M1 Accuracy and Packet Loss Tolerance')
rows=[]
for partition in ['iid','two-class','dirichlet']:
 for method,topo in [('local','ring'),('gossip','ring'),('gossip','grid'),('gossip','random')]:
  g=base(partition,method,topo)[0]
  rows.append([partition,'Local only' if method=='local' else topo.title(),stat(g),stat(g,'final_consensus_error',False),5])
table(['Partition','Method','Accuracy % ± CI','Consensus error ± CI','Seeds'],rows,[1.1,1.1,1.75,2.5,.51])
caption('Table 4. MNIST results at N=5, ten rounds, one local epoch, and no loss or delay. Intervals are 95% across five seeds.')
fig('figures/baseline_accuracy.png','Figure 2. Mean client accuracy versus round, with 95% confidence bands. Conditions share data and initialization within each seed.')
p('IID clients already see all classes, whereas two-class clients cannot learn the entire label space from their own samples alone. The appropriate evidence is the matched communication benefit under each partition. Learning curves reveal intermediate behavior that a final saturated accuracy value can obscure. These results do not establish robustness to impairments that are absent from this table.')

page('4 Results continued')
h('4.1.1 Paired Gains over Local Only')
rows=[]
for r in sorted(summary['paired_vs_local'],key=lambda r: (['iid','two-class','dirichlet'].index(r['partition']),['ring','grid','random'].index(r['topology']))):
 if r.get('dataset','MNIST')!='MNIST' or r.get('rounds',10)!=10:continue
 a=r['accuracy_gain_pp'];rows.append([r['partition'],r['topology'].title(),f"{a['mean']:.2f} ± {(a['ci95_half'] or 0):.2f}",a['n']])
table(['Partition','Topology','Accuracy gain over local only in points','Pairs'],rows,[1.2,1.25,3.75,.76])
caption('Table 5. Within-seed accuracy differences. Each condition has five matched pairs.')
h('4.2 M2 Event Triggered Communication')
rows=[]
for topo in ['ring','grid']:
 for method,delta in [('gossip',0.0),('event',.5),('event',2.0)]:
  g=select(dataset='MNIST',partition='dirichlet',method=method,topology=topo,clients=5,loss=0.0,delay=0,delta=delta)[0]
  rows.append([topo.title(),delta,stat(g),stat(g,'message_savings_pct',False),stat(g,'broadcast_suppression_pct',False)])
table(['Topology','δ','Accuracy % ± CI','Message savings % ± CI','Broadcast suppression % ± CI'],rows,[1.0,.45,1.55,1.9,2.06])
caption('Table 6. Event-triggered Dirichlet MNIST. Actual message savings measure payload reduction; broadcast suppression uses a different denominator.')
fig('figures/event_tradeoff.png','Figure 3. Accuracy and actual message savings versus fixed threshold, with 95% seed intervals.')
ring0=select(dataset='MNIST',partition='dirichlet',method='gossip',topology='ring',clients=5,loss=0.0,delay=0,delta=0.0)[0]
ring2=select(dataset='MNIST',partition='dirichlet',method='event',topology='ring',clients=5,loss=0.0,delay=0,delta=2.0)[0]
p(f"At δ=2.0, Ring mean message savings are {ring2['message_savings_pct']['mean']:.1f}%, while mean accuracy changes from {100*ring0['final_accuracy']['mean']:.2f}% to {100*ring2['final_accuracy']['mean']:.2f}%. These are the observed means for the fixed threshold, not a claim of accuracy-preserving compression.")
p('A silent broadcast can delay the spread of class information. Traffic savings therefore do not imply negligible accuracy loss. An absolute parameter-distance threshold is specific to the model, partition, and training budget; it should not be treated as transferable to another architecture without evaluation.')

page('4.3 Packet Loss Delay and Burst Erasures')
rows=[]
for g in select(dataset='MNIST',partition='dirichlet',method='gossip',clients=5,loss=.2):
 rows.append([g['topology'].title(),g['loss_model'],g['delay'],stat(g),stat(g,'final_consensus_error',False),f"{g['mean_drift']['mean']:.4f}"])
table(['Topology','Channel','Max delay','Accuracy % ± CI','Consensus error ± CI','Mean drift'],rows,[.7,1.15,.6,1.55,1.9,1.06])
caption('Table 7. Impaired-channel results on Dirichlet MNIST. The delay and burst comparisons use Ring; independent zero-delay loss also compares Grid and Random. Drift averages the aggregation-only norm across rounds and seeds. Bernoulli and burst channels match nominal long-run erasure probability, not realized losses.')
bernoulli=select(dataset='MNIST',partition='dirichlet',method='gossip',topology='ring',clients=5,loss=.2,delay=0,loss_model='bernoulli')[0]
burst=select(dataset='MNIST',partition='dirichlet',method='gossip',topology='ring',clients=5,loss=.2,delay=0,loss_model='gilbert-elliott')[0]
if 'realized_drop_pct' in burst:
 p(f"The zero-delay Bernoulli runs erase {bernoulli['realized_drop_pct']['mean']:.1f}% of attempted messages on average, compared with {burst['realized_drop_pct']['mean']:.1f}% for Gilbert-Elliott. Both models have nominal long-run loss 20%; the finite-run realizations need not coincide.")
fig('figures/delay_effect.png' ,'Figure 4. Delay sensitivity with independent erasure probability 0.2. Individual surviving messages can still have zero delay.')
p('These runs exercise the delay model at fixed local-training and gossip-round budgets. An observed change can reflect missing arrivals, stale models, and altered mixing. The newest-version rule prevents an old packet from overwriting a newer accepted model. A favorable outcome at one delay is not evidence that arbitrary latency is harmless.')
p('The burst model starts in stationarity, but ten rounds remain a short observation window relative to its five-round mean bad-state duration. Realized dropped and pending counts must be considered alongside accuracy. Neither this two-state model nor independent erasures constitute a full wireless physical-layer simulation.')
p('Nonzero aggregation-only mean drift is consistent with receiver-wise renormalization under asymmetric reception. It distinguishes a measured limitation of the chosen method from an ideal doubly stochastic averaging property. Replacing this policy with an average-preserving algorithm would be a separate implementation and comparison, rather than a wording correction.')

page('4.4 Scale and CIFAR 10 Extension')
rows=[]
for n in [5,20,50]:
 for topo in ['ring','grid']:
  g=base('two-class','gossip',topo,clients=n)[0]
  rows.append([n,topo.title(),stat(g),stat(g,'final_consensus_error',False),f"{g['attempted_messages']['mean']:.0f}"])
table(['Clients','Topology','Accuracy % ± CI','Consensus error ± CI','Attempts'],rows,[.6,.9,1.75,2.4,1.01])
caption('Table 8. Actual neural-network training at N=5, 20, and 50, using two-class MNIST. Total dataset size is fixed.')
scale5=base('two-class','gossip','ring',clients=5)[0]
scale50=base('two-class','gossip','ring',clients=50)[0]
p(f"Ring mean accuracy is {100*scale5['final_accuracy']['mean']:.2f}% at N=5 and {100*scale50['final_accuracy']['mean']:.2f}% at N=50 under this fixed-round, fixed-total-data protocol. The corresponding uncertainty intervals and traffic counts appear in the scale comparison above.")
p('This scale experiment trains client models; it is not merely a larger eigenvalue calculation. Each client sees two labels, but each label is distributed across more clients as N increases. Total data remain fixed, so per-client examples and minibatch counts decline. Within each N, Ring and Grid share partitions and training budgets. Fixed-density Random graphs are omitted from this training comparison because their growing communication cost would confound a sparse-graph comparison.')
cifar=select(dataset='CIFAR10')
if cifar:
 rows=[]
 for g in sorted(cifar,key=lambda g:g['method']!='local'):rows.append(['Local only' if g['method']=='local' else 'Ring gossip',stat(g),stat(g,'final_consensus_error',False),5])
 table(['CIFAR-10 method','Accuracy % ± CI','Consensus error ± CI','Seeds'],rows,[1.7,1.8,2.7,.76])
 caption('Table 9. CIFAR-10 with five clients, Dirichlet α=0.1, and the same ten-round budget.')
 p('CIFAR-10 contains 50,000 training and 10,000 test RGB images, each 32×32. Channel means are (0.4914, 0.4822, 0.4465), and standard deviations are (0.2470, 0.2435, 0.2616). The two-convolution model has three input channels and a 2048→128 first fully connected layer, totaling 268,650 parameters. This extension compares local-only and Ring gossip; it does not claim a complete CIFAR-10 topology or channel sweep. No augmentation or architecture search is performed.')
 cifar_local=[g for g in cifar if g['method']=='local'][0]
 cifar_ring=[g for g in cifar if g['method']=='gossip'][0]
 p(f"Under this protocol, Ring mean accuracy exceeds local-only by {100*(cifar_ring['final_accuracy']['mean']-cifar_local['final_accuracy']['mean']):.2f} percentage points. The table reports seed uncertainty for both means; the paired seed-level difference is also retained in the accompanying summary JSON.")
 p('The short-budget small-CNN result is a controlled comparison rather than a competitive CIFAR-10 benchmark. Absolute accuracy and any paired gain must be interpreted within this restricted protocol. Test outcomes are reported without retrospectively changing the budget to obtain a preferred conclusion.')
else:
 p('CIFAR-10 support is implemented, but no completed five-seed CIFAR-10 result is included here. Empirical claims remain restricted to MNIST; harder-task conclusions require completing the extension.')

page('4.5 M3 Spectral Gap and Convergence Analysis')
h('Spectral Gap')
p('The ideal spectral gap γ=1−ρ uses the second-largest eigenvalue modulus ρ. The signed spectrum and the magnitude used for SLEM are retained separately.')
fig('figures/topologies.png','Figure 5. Actual five-node graphs at seed 42. The trimmed Grid has degree sequence 2, 3, 1, 2, 2, including a pendant node.')
rows=[]
for s in spectra:rows.append([s['topology'].title(),s['n_edges'],f"{s['spectral_gap']:.6f}",f"{s['slem']:.6f}",int(s['mixing_time_1e2']),int(s['mixing_time_1e3'])])
table(['Topology','Edges','Gap γ','SLEM ρ','T(.01)','T(.001)'],rows,[1.2,.6,1.3,1.3,1.28,1.28])
caption('Table 10. Ideal five-node spectra. T is a sufficient integer number of rounds for a relative squared-disagreement target.')
h('Eigenvalue Spectra')
rows=[[s['topology'].title()]+[f'{0.0 if abs(v)<1e-12 else v:.6f}' for v in s['all_eigenvalues']] for s in spectra]
table(['Topology','λ₁','λ₂','λ₃','λ₄','λ₅'],rows,[1.05,1.18,1.18,1.18,1.18,1.19])
caption('Table 11. Signed eigenvalues sorted by descending magnitude. SLEM uses magnitude; the full spectrum preserves signs.')
h('Convergence Bound')
p('For fixed symmetric doubly stochastic W, let J=11ᵀ/N and ρ=||W−J||₂. Disagreement under pure averaging is (W−J)ᵗX(0). Applying the spectral-norm bound to squared Frobenius disagreement gives [3,4]:')
eq('E(t) ≤ E(0) ρ²ᵗ,     γ = 1 − ρ',8)
h('Mixing Time')
eq('T(ε) = ceiling{ log(1/ε) / [2 log(1/ρ)] }',9)
p('Equation (9) targets E(t)/E(0)≤ε. An unsquared-norm target would remove the factor of two in the denominator. The table therefore states the relative squared-disagreement convention and reports sufficient integer round counts.')

page('4.6 Spectral Validation and Comparison with M1')
fig('figures/spectral_bound.png','Figure 6. Scalar averaging variance and the squared-error bound on identical graph instances. This checks ideal averaging, not SGD convergence.')
p('A scalar trajectory lies below the bound up to numerical effects. It need not equal the bound because the initial vector may have little projection onto the slowest eigenmode. Near floating-point precision, the apparent decay slope should not be interpreted as a physical convergence rate. The scalar trajectory is interpreted as a check of the stated upper bound rather than evidence of an exact rate match.')
h('Cross Comparison with M1')
p('Local SGD continually injects disagreement. Loss, event suppression, and stale arrivals also change the update operator with time. The fixed-matrix proof does not cover these effects. A larger ideal spectral gap consequently does not guarantee higher trained-model accuracy or smaller final learning disagreement.')
p('The original saved results already contradict an unconditional ranking: under 20% loss, Random has consensus error about 0.161 and Ring about 0.301, although the ideal Ring gap is larger. The revised study treats these training outcomes as empirical comparisons rather than confirmations of an unconditional spectral prediction.')
fig('m3/m3_spectral_gap_vs_size.png','Figure 7. Ideal gap and squared-error mixing bound at N=5, 20, and 50. Fixed p=0.5 Random graphs become denser as N increases.')

discussion(p,h,page,select,base,stat,summary)

page('6 Conclusion')
p(f"On two-class MNIST, Ring gossip changes mean client accuracy from {100*two['final_accuracy']['mean']:.2f}% under local-only training to {100*ring['final_accuracy']['mean']:.2f}%. On IID MNIST, the local-only control already reaches {100*iid['final_accuracy']['mean']:.2f}%. The contrast makes data heterogeneity central to interpreting the value of communication.")
p('This study evaluates established decentralized learning mechanisms using matched local-only controls, heterogeneous data, five-seed uncertainty, and corrected traffic accounting. It separates trained-model outcomes from the ideal properties of a fixed mixing matrix. Conclusions apply to the tested partitions, budgets, graph instances, and channel models; they do not establish universal packet-loss immunity or a universal Ring-over-Grid preference.')
p('Event triggering is an empirical accuracy-traffic tradeoff measured with actual directed-message and payload-byte counts. Delay is evaluated using version-aware receivers and a finite horizon. Saved runs, scripts, and dependency specifications make the numerical tables reproducible, while a CPU checkpoint demonstration supports review on another laptop.')
h('6.1 Future Work')
p('Future evaluation should compare selected unreliable-link and communication-efficient algorithms under the same data, training, and byte budgets. Useful extensions include average-preserving aggregation under directed loss, time-varying graphs, validation-driven threshold selection, longer harder-task training, and client fairness. Privacy protection and vehicular mobility require dedicated mechanisms rather than an inference from server-free execution {ref:r12}{ref:r19}{ref:r30}.')
doc.add_heading('7 Declarations',1)
p('Funding: No specific grant is declared for this project. Conflicts of interest: None declared by the project authors. Data availability: MNIST and CIFAR-10 are public benchmark datasets [9,10]. Code and results: The accompanying project includes source, dependency specifications, raw results, statistics, figures, validation tests, and a portable demonstration checkpoint. The accompanying artifacts provide the evaluated implementation and supporting evidence.')
page('8 References')
missing=set(references)-set(reference_numbers)
if missing:raise RuntimeError('Uncited bibliography entries: '+str(missing))
for key,number in reference_numbers.items():
 r={k:html.unescape(v) if isinstance(v,str) else v for k,v in references[key].items()}
 year=str(r['year']) if r.get('year') else 'n.d.'
 venue=r['venue']
 if r.get('volume'):
  venue+=', '+str(r['volume'])
  if r.get('issue'):venue+='('+str(r['issue'])+')'
 if r.get('pages'):venue+=', '+str(r['pages'])
 text=f"[{number}] {r['authors']} ({year}). {r['title']}. {venue}. {r['url']}"
 para=doc.add_paragraph(text)
 para.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 para.paragraph_format.space_after=Pt(8)
 para.paragraph_format.keep_together=True
 for run in para.runs:run.font.size=Pt(9.5)
(ROOT/'report_sources/reference_numbering.json').write_text(json.dumps(reference_numbers,indent=2),encoding='utf-8')
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
path=(ROOT/'.work/report_preview.docx') if '--preview' in sys.argv else (OUT/'Team21_Revised_Decentralized_Learning_Report.docx');doc.save(path)
print(path);print('Completed runs used:',summary['completed_runs'])
